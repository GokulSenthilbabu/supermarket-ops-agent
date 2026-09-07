from pathlib import Path
from datetime import date, datetime, time, timedelta
from decimal import Decimal

import matplotlib

matplotlib.use("Agg")

import matplotlib.pyplot as plt

from pptx import Presentation
from pptx.util import Inches, Pt

from sqlalchemy import select, func

from app.database import SessionLocal
from app.models import Bill, BillItem, Product


BASE_DIR = Path(__file__).resolve().parent.parent.parent
REPORT_DIR = BASE_DIR / "generated_reports"

REPORT_DIR.mkdir(exist_ok=True)


def get_weekly_data():
    db = SessionLocal()

    try:
        today = date.today()

        week_start = today - timedelta(days=6)

        start_datetime = datetime.combine(
            week_start,
            time.min,
        )

        end_datetime = datetime.combine(
            today,
            time.max,
        )

        bills = db.execute(
            select(Bill).where(
                Bill.status == "FINALIZED",
                Bill.finalized_at >= start_datetime,
                Bill.finalized_at <= end_datetime,
            )
        ).scalars().all()

        total_sales = Decimal("0")
        total_tax = Decimal("0")

        cash = Decimal("0")
        upi = Decimal("0")
        card = Decimal("0")

        daily_sales = {}

        for i in range(7):
            current_day = week_start + timedelta(days=i)
            daily_sales[current_day.isoformat()] = Decimal("0")

        for bill in bills:

            amount = Decimal(
                str(bill.grand_total or 0)
            )

            tax = (
                Decimal(str(bill.cgst or 0))
                + Decimal(str(bill.sgst or 0))
                + Decimal(str(bill.igst or 0))
            )

            total_sales += amount
            total_tax += tax

            mode = (bill.payment_mode or "").upper()

            if mode == "CASH":
                cash += amount

            elif mode == "UPI":
                upi += amount

            elif mode == "CARD":
                card += amount

            if bill.finalized_at:
                bill_date = bill.finalized_at.date()

                if bill_date.isoformat() in daily_sales:
                    daily_sales[
                        bill_date.isoformat()
                    ] += amount

        # ---------------------------------------------
        # TOP PRODUCTS
        # ---------------------------------------------

        rows = db.execute(
            select(
                Product.name,
                func.sum(BillItem.quantity).label(
                    "quantity"
                ),
                func.sum(BillItem.total).label(
                    "sales"
                ),
            )
            .join(
                BillItem,
                BillItem.product_id == Product.id,
            )
            .join(
                Bill,
                Bill.id == BillItem.bill_id,
            )
            .where(
                Bill.status == "FINALIZED",
                Bill.finalized_at >= start_datetime,
                Bill.finalized_at <= end_datetime,
            )
            .group_by(
                Product.id,
                Product.name,
            )
            .order_by(
                func.sum(BillItem.quantity).desc()
            )
            .limit(10)
        ).all()

        top_products = []

        for row in rows:
            top_products.append({
                "name": row.name,
                "quantity": float(row.quantity or 0),
                "sales": float(row.sales or 0),
            })

        return {
            "success": True,
            "week_start": week_start.isoformat(),
            "week_end": today.isoformat(),
            "bill_count": len(bills),
            "total_sales": float(total_sales),
            "total_tax": float(total_tax),
            "cash": float(cash),
            "upi": float(upi),
            "card": float(card),
            "daily_sales": {
                key: float(value)
                for key, value in daily_sales.items()
            },
            "top_products": top_products,
        }

    except Exception as exc:
        return {
            "success": False,
            "error": str(exc),
        }

    finally:
        db.close()


def create_sales_chart(data, path):
    dates = list(data["daily_sales"].keys())
    values = list(data["daily_sales"].values())

    labels = [
        datetime.strptime(
            value,
            "%Y-%m-%d"
        ).strftime("%d %b")
        for value in dates
    ]

    plt.figure(figsize=(10, 5))

    plt.plot(
        labels,
        values,
        marker="o",
    )

    plt.title("Daily Sales")
    plt.xlabel("Date")
    plt.ylabel("Sales (INR)")

    plt.xticks(rotation=45)

    plt.tight_layout()

    plt.savefig(path, dpi=150)

    plt.close()


def create_top_products_chart(data, path):
    products = data["top_products"][:5]

    names = [
        product["name"]
        for product in products
    ]

    quantities = [
        product["quantity"]
        for product in products
    ]

    plt.figure(figsize=(10, 5))

    plt.bar(
        names,
        quantities,
    )

    plt.title("Top Products by Quantity Sold")
    plt.xlabel("Product")
    plt.ylabel("Quantity")

    plt.xticks(
        rotation=30,
        ha="right",
    )

    plt.tight_layout()

    plt.savefig(path, dpi=150)

    plt.close()


def create_payment_chart(data, path):
    labels = [
        "Cash",
        "UPI",
        "Card",
    ]

    values = [
        data["cash"],
        data["upi"],
        data["card"],
    ]

    plt.figure(figsize=(7, 7))

    plt.pie(
        values,
        labels=labels,
        autopct="%1.1f%%",
    )

    plt.title("Payment Method Distribution")

    plt.tight_layout()

    plt.savefig(path, dpi=150)

    plt.close()


def generate_weekly_sales_deck():
    data = get_weekly_data()

    if not data["success"]:
        return data

    sales_chart = REPORT_DIR / "weekly_sales_chart.png"
    products_chart = REPORT_DIR / "top_products_chart.png"
    payment_chart = REPORT_DIR / "payment_chart.png"

    create_sales_chart(
        data,
        sales_chart,
    )

    create_top_products_chart(
        data,
        products_chart,
    )

    create_payment_chart(
        data,
        payment_chart,
    )

    presentation = Presentation()

    # ---------------------------------------------
    # SLIDE 1
    # ---------------------------------------------

    slide = presentation.slides.add_slide(
        presentation.slide_layouts[0]
    )

    slide.shapes.title.text = (
        "Weekly Sales Analysis"
    )

    slide.placeholders[1].text = (
        f"{data['week_start']} to "
        f"{data['week_end']}"
    )

    # ---------------------------------------------
    # SLIDE 2
    # ---------------------------------------------

    slide = presentation.slides.add_slide(
        presentation.slide_layouts[5]
    )

    slide.shapes.title.text = "Sales Overview"

    textbox = slide.shapes.add_textbox(
        Inches(1),
        Inches(1.5),
        Inches(8),
        Inches(4),
    )

    frame = textbox.text_frame

    lines = [
        f"Total Sales: ₹{data['total_sales']:,.2f}",
        f"GST Collected: ₹{data['total_tax']:,.2f}",
        f"Number of Bills: {data['bill_count']}",
        f"Cash: ₹{data['cash']:,.2f}",
        f"UPI: ₹{data['upi']:,.2f}",
        f"Card: ₹{data['card']:,.2f}",
    ]

    for index, line in enumerate(lines):

        paragraph = (
            frame.paragraphs[0]
            if index == 0
            else frame.add_paragraph()
        )

        paragraph.text = line
        paragraph.font.size = Pt(22)

    # ---------------------------------------------
    # SLIDE 3
    # ---------------------------------------------

    slide = presentation.slides.add_slide(
        presentation.slide_layouts[5]
    )

    slide.shapes.title.text = "Daily Sales"

    slide.shapes.add_picture(
        str(sales_chart),
        Inches(1),
        Inches(1.3),
        width=Inches(8),
    )

    # ---------------------------------------------
    # SLIDE 4
    # ---------------------------------------------

    slide = presentation.slides.add_slide(
        presentation.slide_layouts[5]
    )

    slide.shapes.title.text = "Top Products"

    slide.shapes.add_picture(
        str(products_chart),
        Inches(1),
        Inches(1.3),
        width=Inches(8),
    )

    # ---------------------------------------------
    # SLIDE 5
    # ---------------------------------------------

    slide = presentation.slides.add_slide(
        presentation.slide_layouts[5]
    )

    slide.shapes.title.text = "Payment Methods"

    slide.shapes.add_picture(
        str(payment_chart),
        Inches(2),
        Inches(1.3),
        width=Inches(6),
    )

    # ---------------------------------------------
    # SLIDE 6
    # ---------------------------------------------

    slide = presentation.slides.add_slide(
        presentation.slide_layouts[5]
    )

    slide.shapes.title.text = "Business Insights"

    textbox = slide.shapes.add_textbox(
        Inches(1),
        Inches(1.5),
        Inches(8),
        Inches(4),
    )

    frame = textbox.text_frame

    insights = []

    if data["top_products"]:
        top = data["top_products"][0]

        insights.append(
            f"Top-selling item: {top['name']} "
            f"({top['quantity']:.0f} units)"
        )

    if data["total_sales"] > 0:
        insights.append(
            f"Average bill value: "
            f"₹{data['total_sales'] / max(data['bill_count'], 1):,.2f}"
        )

    if data["upi"] > data["cash"] and data["upi"] > data["card"]:
        insights.append(
            "UPI is the dominant payment method."
        )

    if data["total_tax"] > 0:
        insights.append(
            f"Total GST collected: "
            f"₹{data['total_tax']:,.2f}"
        )

    if not insights:
        insights.append(
            "No finalized sales were recorded during this period."
        )

    for index, insight in enumerate(insights):

        paragraph = (
            frame.paragraphs[0]
            if index == 0
            else frame.add_paragraph()
        )

        paragraph.text = insight
        paragraph.font.size = Pt(20)

    # ---------------------------------------------
    # SAVE
    # ---------------------------------------------

    output_path = (
        REPORT_DIR
        / "weekly_sales_analysis.pptx"
    )

    presentation.save(output_path)

    return {
        "success": True,
        "pptx_path": str(output_path),
        "week_start": data["week_start"],
        "week_end": data["week_end"],
        "total_sales": data["total_sales"],
        "total_tax": data["total_tax"],
        "bill_count": data["bill_count"],
    }